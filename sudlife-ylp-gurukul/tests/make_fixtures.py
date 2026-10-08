"""Generate fictional test files for every dataset type the application imports.

Usage: python3 tests/make_fixtures.py <output-folder>
All names, IDs and content are invented. Some rows are deliberately wrong so
the validation rules can be tested.
"""
import json, sys, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font

out = sys.argv[1]
os.makedirs(out, exist_ok=True)


def csv(name, header, rows):
    import csv as c
    with open(os.path.join(out, name), "w", newline="", encoding="utf-8") as f:
        w = c.writer(f)
        w.writerow(header)
        w.writerows(rows)


# ---------- 1. Participant master (header on row 3, Excel dates, ambiguous text date)
PEOPLE = [
    (100234, "Kavya Menon", dt.date(2026, 4, 6), "Band A2", "Executive - Learning & Development", "Learning & Development", "Human Resources", "FY2026"),
    (100235, "Rahul Verma", dt.date(2025, 11, 12), "Band B1", "Assistant Manager - FP&A", "Financial Planning", "Finance", "FY2025"),
    (100236, "Ananya Ghosh", "03/04/2026", "Band A2", "Executive", "Underwriting", "Operations", "Batch 2026"),
    (100237, "Farhan Siddiqui", "14/05/2026", "Band A1", "Engineer - Digital", "Digital Platforms", "Technology", "FY2026"),
    (100238, "Lakshmi Pillai", dt.date(2026, 6, 1), "Band A1", "Officer", "Compliance", "Risk & Compliance", "FY 2025-26"),
    (100239, "Neel Joshi", dt.date(2026, 7, 15), "", "Pricing Analyst", "Pricing", "Actuarial", "FY2026"),
    (100240, "Ritu Saxena", dt.date(2025, 12, 4), "Band B1", '<img src=x onerror="window.__xss=1">', "Customer Service", "Customer Experience", "FY2025"),
    (100241, "Omkar Patil", dt.date(2026, 5, 20), "Band A2", "Executive", "Bancassurance Support", "Distribution", "FY2026"),
]
wb = Workbook(); ws = wb.active; ws.title = "Participant Master"
ws["A1"] = "YLP Participant Master - Confidential"; ws["A1"].font = Font(bold=True, size=14)
ws.append([]); ws.append(["Employee Id", "Status", "Full Name", "Date Of Joining", "Job Level", "Designation", "Current Department", "Business Unit", "Direct Manager Email Id", "HRBP Email Id", "Official Email Id", "YLP Batch"])
for i, p in enumerate(PEOPLE):
    ws.append([p[0], "Active", p[1], p[2], p[3], p[4], p[5], p[6], f"mgr{i}@example.com", "hrbp@example.com", f"person{i}@example.com", p[7]])
for row in ws.iter_rows(min_row=4, min_col=4, max_col=4):
    for c in row:
        if isinstance(c.value, dt.date): c.number_format = "DD-MMM-YYYY"
wb.save(f"{out}/YLP_Participant_Master.xlsx")

# ---------- 2. Previous FY KRA collated workbook (spec headers; mixed weight formats)
H = ["Employee ID", "Employee Name", "Business Unit", "Department", "Goals / Key Result Areas KRA", "Goals / Key Result Areas Description", "Goals / Key Result Areas Weightage Percentage"]
wb = Workbook(); ws = wb.active; ws.title = "Employee_KRAs"; ws.append(H)
for r in [
    [100235, "Rahul Verma", "Finance", "Financial Planning", "Budget Variance Reporting", "Monthly variance packs", 0.4],
    [100235, "Rahul Verma", "Finance", "Financial Planning", "Cost Allocation Accuracy", "Allocate shared costs", 0.35],
    [100235, "Rahul Verma", "Finance", "Financial Planning", "Process Improvement", "Maintain reconciliation MIS", 0.25],
    [100240, "Ritu Saxena", "Customer Experience", "Customer Service", "Complaint Resolution", "Resolve within TAT", "50%"],
    [100240, "Ritu Saxena", "Customer Experience", "Customer Service", "Complaint Resolution", "Escalations within TAT", "45%"],
    [999001, "Not In Master", "Finance", "Investments", "Portfolio Reporting", "Weekly reports", 1],
]: ws.append(r)
wb.create_sheet("Notes").append(["Not a data sheet"])
wb.save(f"{out}/YLP_KRA_Collated_FY25_26.xlsx")

# ---------- 3. Current FY KRA workbook with pandas-merge columns
wb = Workbook(); ws = wb.active; ws.title = "Employee_KRAs"
ws.append(["Employee ID", "Employee Name", "Business Unit_x", "Department", "Goals / Key Result Areas KRA", "Goals / Key Result Areas Description", "Goals / Key Result Areas Weightage Percentage", "Business Unit_y"])
for r in [
    [100235, "Rahul Verma", "Finance", "Financial Planning", "Automate first-premium reconciliation", "Implement automated matching", 0.5, "Finance"],
    [100235, "Rahul Verma", "Finance", "Financial Planning", "Lead the cross-functional cost project", "Own and drive the cost project", 0.5, "Finance"],
    [100237, "Farhan Siddiqui", "Technology", "Digital Platforms", "Deliver platform releases", "Deliver releases on the sprint plan", 0.6, "Technology"],
    [100237, "Farhan Siddiqui", "Technology", "Digital Platforms", "Release quality", "Defect leakage", 0.395, "Technology"],
]: ws.append(r)
wb.save(f"{out}/YLP_KRA_FY26_27_Current.xlsx")

# ---------- 4. Individual goal sheet (KRA Dashboard + Submitted Metric Scales)
GOAL = [("Learning Platforms Management", 0.2), ("Learning Operations and Administration", 0.2), ("Department-Wise Learning Accountability", 0.2),
        ("Learning Content, Instructional Design and Communication", 0.2), ("Compliance Training Calendar", 0.15), ("Personal Development and Capability Building", 0.05)]
wb = Workbook(); ws = wb.active; ws.title = "KRA Dashboard"
ws["A1"] = "KRA Goal Sheet - FY 2026-27 (Revised)"; ws.merge_cells("A1:F1")
ws["A3"] = "Employee Name:"; ws["B3"] = "Kavya Menon"
ws["A4"] = "Designation:"; ws["B4"] = "Executive - Learning & Development"
ws["A5"] = "Financial Year:"; ws["B5"] = "FY 2026-27"
ws.append([]); ws.append(["KRA No", "KRA Name", "Weightage", "SMART Goal Statement", "Key Parameters", "KRA Type"])
for i, (name, w) in enumerate(GOAL, 1):
    ws.append([i, name, w, f"Deliver {name.lower()} outcomes by Q4", "Timeliness; Quality", "Functional" if i < 6 else "Development"])
    ws.cell(row=ws.max_row, column=3).number_format = "0%"
ws.append(["Total", "", 1.0])
sc = wb.create_sheet("Submitted Metric Scales")
sc["A1"] = "Submitted Metric Scales"; sc.append([])
sc.append(["KRA", "Weight", "Parameter", "Level 1", "Level 2", "Level 3", "Level 4", "Level 5", "Level 5+", "Rationale and Justification"])
r = 4
for name, w in GOAL[:5] + [("Personal Development & Capability", 0.05)]:
    for p in ("Timeliness", "Quality"):
        sc.append([name if p == "Timeliness" else None, w if p == "Timeliness" else None, p, "<60%", "60-74%", "75-89%", "90-99%", "100%", ">100%", "Agreed with manager"])
    sc.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
    sc.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)
    r += 2
wb.save(f"{out}/Kavya_Menon_KRA_GoalSheet_FY2026-27_Revised.xlsx")

# ---------- 5. CSV with BOM, quoted comma, newline, script and formula injection
open(f"{out}/participants_injection.csv", "w", encoding="utf-8", newline="").write(
    "﻿Employee ID,Employee Name,Date Of Joining,Job Level,Designation,Department,Business Unit,YLP Batch\r\n"
    '200001,"Sharma, Priya",2026-04-06,Band A1,"Officer\nSecond line",Claims,Operations,FY2026\r\n'
    '200002,"<script>window.__xss=2</script>",14-Apr-2026,Band A1,Officer,Claims,Operations,FY2026\r\n'
    '200003,"=HYPERLINK(""http://evil.example"",""click"")",2026-04-20,Band A2,Executive,Legal,Risk & Compliance,FY2026\r\n'
    '200004,"Dev ""DJ"" Jain",2026-05-01,Band A2,Executive,Legal,Risk & Compliance,FY2026\r\n')

# ---------- 6. JSON array with a __proto__ key
json.dump([
    {"Employee ID": "300001", "Employee Name": "Json One", "Date Of Joining": "2026-04-01", "Job Level": "Band A1", "Department": "Audit", "Business Unit": "Risk & Compliance", "YLP Batch": "FY2026"},
    {"Employee ID": "300002", "Employee Name": "Json Two", "Date Of Joining": "2025-11-01", "Job Level": "Band A2", "Department": "Audit", "Business Unit": "Risk & Compliance", "YLP Batch": "FY2025", "__proto__": {"polluted": True}},
], open(f"{out}/participants.json", "w"))

# ---------- 7. Role descriptor workbook (HRMS headers)
wb = Workbook(); ws = wb.active; ws.title = "Role Descriptors"
ws.append(["Business Unit (as per HRMS)", "Department (as per HRMS)", "Work Role", "Key descriptors or deliverables", "Distributed / COE", "WTW Title", "WTW Job Role", "Unique Role Name", "Special Skills / Qualifications", "Role (IC/TM)"])
ws.append(["Human Resources", "Learning & Development", "Executive - Learning & Development", "Learning platforms, compliance training calendar, content", "COE", "Executive - Learning & Development", "L&D Executive", "L&D Executive", "LMS", "IC"])
ws.append(["Finance", "Financial Planning", "Assistant Manager - FP&A", "Variance reports, reconciliation, cost allocation, budget", "COE", "Assistant Manager - FP&A", "Finance Analyst", "FP&A Assistant Manager", "Excel", "IC"])
ws.append(["Technology", "Digital Platforms", "Engineer - Digital", "Platform releases and automation", "COE", "Engineer - Digital", "Software Engineer", "Digital Platform Engineer", "Agile", "IC"])
ws.append(["Operations", "Underwriting", "Underwriter", "Underwrite proposals", "COE", "Underwriter", "Underwriter", "", "Underwriting", "IC"])  # missing unique role
wb.save(f"{out}/YLP_Role_Descriptors.xlsx")

# ---------- 8. Leadership competency workbook (detailed + concise sheets)
DOMAINS = {
    "Customer & Stakeholder Management": ["Communication and Networking", "Influencing and Negotiation", "Customer Engagement and Service Orientation", "Emotional Intelligence", "Collaboration"],
    "Performance and Result Orientation": ["Achievement Orientation and Ownership of Outcomes", "Continuous Improvement", "Creating High-Performance Teams and Talent Development"],
    "Strategy and Innovation": ["Business Acumen and Strategic Planning", "Customer-Centric Innovation", "Problem Solving and Decision Making"],
    "Risk Management": ["Risk Assessment and Compliance", "Change Planning and Implementation", "Governance and Monitoring", "Ethics"],
    "Digital Mindset and Agility": ["Digital Mindset", "Data Analytics Proficiency", "Agile Working"],
}
HDR = ["Competency", "Behavioral Anchors", "Definitions of the behavioral anchors", "Level 1: Basic", "Level 2: Developing", "Level 3: Proficient", "Level 4: Advanced", "Level 5: Expert"]
wb = Workbook(); d = wb.active; d.title = "Detailed Master"; d.append(HDR)
c2 = wb.create_sheet("Concise"); c2.append(HDR)
for dom, anchors in DOMAINS.items():
    for i, a in enumerate(anchors):
        d.append([dom if i == 0 else None, a, f"IMPORTED definition of {a}: what good looks like across SUD Life teams and stakeholders.",
                  f"Basic {a}: applies with guidance in routine work.", f"Developing {a}: applies independently in routine work.",
                  f"Proficient {a}: applies consistently, including non-routine situations, with results.", f"Advanced {a}: applies in complex situations and coaches others.",
                  f"Expert {a}: shapes practice across the organisation."])
        c2.append([dom if i == 0 else None, a, a, "Basic", "Developing", f"Concise L3 for {a}", "Advanced", "Expert"])
wb.save(f"{out}/Leadership_Competency_Framework.xlsx")

# ---------- 9. Competency assessments CSV (validated, self, bad rows)
rows = []
anchors = [a for lst in DOMAINS.values() for a in lst]
for a in anchors:
    rows.append(["100235", "2026-09-20", "Manager", "", a, 2, 3, 3, "Anita Rao", "Two examples including a non-routine escalation", 2, "Yes", "Manager", "Validated", "", "2026-12-20", "Current"])
    rows.append(["100236", "2026-09-20", "Self", "", a, "", 3, 3, "Ananya Ghosh (self)", "", 0, "No", "", "", "", "2026-12-20", "Self-assessed"])
rows.append(["100235", "2026-09-20", "Manager", "", "Ethics", "", 7, 3, "Anita Rao", "x", 2, "Yes", "Manager", "Validated", "", "", "Current"])            # invalid score
rows.append(["100235", "2026-09-21", "Manager", "", "Ethics", "", "Meets Expectations", 3, "Anita Rao", "x", 2, "Yes", "Manager", "Validated", "", "", ""])  # KRA rating confusion
rows.append(["100235", "2026-09-22", "Sponsor", "", "Collaboration", "", 3, 4, "Sponsor A", "Panel", 2, "Yes", "Sponsor", "Validated", "", "", ""])        # target not 3
rows.append(["100235", "2026-09-22", "Manager", "", "Telepathy", "", 3, 3, "Anita Rao", "x", 2, "Yes", "Manager", "Validated", "", "", ""])              # unknown anchor
csv("competency_assessments.csv", ["Employee ID", "Assessment Date", "Assessment Type", "Competency Domain", "Behavioural Anchor", "Baseline Level", "Current Level", "Target Level", "Assessor", "Evidence", "Evidence Count", "Non-Routine Evidence", "Validator Role", "Validation Status", "Development Action", "Next Review Date", "Assessment Status"], rows)

# ---------- 10. Projects + project participants
csv("projects.csv", ["ProjectID", "ProjectTitle", "ProjectTheme", "BusinessProblem", "ProjectSponsor", "YLPAdviser", "FunctionsRepresented", "StartDate", "TargetDate", "Status", "CompletionPercentage", "BaselineMeasure", "TargetImpact", "ActualImpact", "PanelScore", "ExceptionApproved"], [
    ["ACT-P1", "Faster underwriting decisions", "Operational efficiency", "Proposals wait too long", "Head of Operations", "Adviser A", "Operations, Finance, Technology", "2026-06-01", "2026-12-31", "In Progress", 45, "9 days", "5 days", "", "", "No"],
    ["ACT-P2", "Learning nudges", "Risk and governance", "Low completion", "CCO", "", "HR, Technology", "2026-06-01", "2026-12-31", "In Progress", 150, "70%", "95%", "", "", "No"],
])
csv("project_participants.csv", ["ProjectID", "EmployeeID", "ParticipantRole", "IndividualContribution", "StakeholderFeedback", "SponsorFeedback", "ContributionRating", "Evidence"], [
    ["ACT-P1", "100235", "Lead", "Mapped the decision flow", "", "", 4, "Flow map"],
    ["ACT-P1", "100236", "Analysis", "", "", "", "", ""],
    ["ACT-P1", "100237", "Automation", "Built the triage script", "", "", 3, "Demo"],
    ["ACT-P9", "100240", "Member", "", "", "", "", ""],   # unknown project
])

# ---------- 11. Other employee-linked tables
csv("development_actions.csv", ["ActionID", "EmployeeID", "DevelopmentCategory", "DevelopmentAction", "CompetencyLinked", "Owner", "StartDate", "TargetDate", "Status", "CompletionPercentage", "Evidence", "ManagerComment", "NextReviewDate"], [
    ["A-1", "100235", "Experience", "Lead the cost project", "Collaboration", "Rahul", "2026-06-01", "2026-09-30", "In Progress", 60, "", "", "2026-10-30"],
    ["A-2", "100235", "Learning", "Finish analytics module", "Data Analytics Proficiency", "Rahul", "2026-06-01", "", "Not Started", 0, "", "", ""],
])
csv("learning_records.csv", ["EmployeeID", "LearningTitle", "LearningType", "CompletionDate", "LearningHours", "Status", "AssessmentScore", "ApplicationEvidence"], [
    ["100235", "SUD Life business model", "SUD Life Business Learning", "2026-08-01", 4, "Completed", 80, "Applied to the cost project"],
    ["100235", "Insurance value chain foundations", "Insurance Value Chain", "2026-08-15", 6, "Completed", 55, ""],
])
csv("mentoring_records.csv", ["EmployeeID", "InteractionType", "Mentor", "MeetingDate", "DiscussionTheme", "ActionCommitted", "ActionCompleted", "NextReviewDate", "Evidence"], [
    ["100235", "Mentoring", "Mentor X", "2026-08-10", "Stakeholders", "Map stakeholders", "Yes", "2026-11-10", "Log 1"],
    ["100235", "", "Mentor X", "2026-09-10", "Prioritising", "", "No", "", ""],
])
csv("stakeholder_feedback.csv", ["FeedbackID", "EmployeeID", "FeedbackDate", "FeedbackProvider", "ProviderRole", "Context", "CompetencyLinked", "FeedbackSummary", "Rating", "Evidence"], [
    ["F-1", "100235", "2026-09-01", "Head of Operations", "Project Sponsor", "ACT-P1", "Collaboration", "Strong facilitation", 4, "Note"],
])
csv("journey_activities.csv", ["ActivityID", "EmployeeID", "JourneyStage", "ActivityName", "ActivityType", "StartDate", "TargetDate", "CompletionDate", "Status", "CompletionPercentage", "Evidence", "Owner"], [
    ["J-1", "100235", "Abhyas", "Stretch: own variance commentary", "Stretch Assignment", "2026-05-01", "2026-08-31", "2026-08-30", "Completed", 100, "", "Manager"],
    ["J-2", "100235", "", "Shadow the CFO review", "Leadership Shadowing", "2026-09-01", "2026-09-30", "", "In Progress", 50, "", "Manager"],
    ["J-3", "100235", "Abhyas", "Mystery task", "Unknown Type", "2026-09-01", "2026-09-30", "", "In Progress", 120, "", "Manager"],
])
# Rating scales by name only (no Employee ID): needs name matching
csv("kra_rating_scales_by_name.csv", ["Employee Name", "Financial Year", "KRA", "Weight", "Parameter", "Level 1", "Level 2", "Level 3", "Level 4", "Level 5", "Level 5+", "Rationale and Justification"], [
    ["Rahul Verma", "FY 2026-27", "Automate first-premium reconciliation", 0.5, "Unmatched share", ">10%", "8-10%", "5-8%", "4-5%", "<4%", "<2%", "Agreed"],
])

# ---------- 12. Complete workbook (all tables + one unrecognised sheet)
wb = Workbook(); first = True
def sheet(name, header, rows):
    global first
    ws = wb.active if first else wb.create_sheet(name)
    if first: ws.title = name; first = False
    ws.append(header)
    for r in rows: ws.append(r)
sheet("Participants", ["EmployeeID", "EmployeeName", "Status", "Cohort", "DateOfJoining", "JobLevel", "Designation", "BusinessUnit", "Department", "JourneyStage", "LeadershipPathway", "MentorStatus", "CareerAspiration"], [
    ["W-01", "Workbook One", "Active", "FY2026", "2026-04-15", "Band A1", "Officer", "Operations", "Claims", "Abhyas", "Project Leadership", "Allocated", "Lead claims"],
    ["W-02", "Workbook Two", "Active", "FY2026", "2026-05-15", "Band A2", "Executive", "Finance", "Treasury", "Bodh", "Specialist Leadership", "Pending", ""],
    ["W-03", "Workbook Three", "Active", "FY2025", "2025-12-15", "Band B1", "Assistant Manager", "Technology", "Data", "Margdarshan", "People Leadership", "Allocated", "Lead a team"],
])
sheet("KRAs", ["EmployeeID", "EmployeeName", "FinancialYear", "KRASequence", "KRAName", "KRADescription", "Weightage"], [
    ["W-01", "Workbook One", "FY 2026-27", 1, "Claims TAT", "Within TAT", 0.6], ["W-01", "Workbook One", "FY 2026-27", 2, "Quality", "Audit", 0.4]])
sheet("Competencies", ["EmployeeID", "AssessmentDate", "AssessmentType", "CompetencyDomain", "BehaviouralAnchor", "CurrentLevel", "Assessor", "Evidence", "EvidenceCount", "NonRoutineEvidence", "ValidatorRole", "ValidationStatus"], [
    ["W-01", "2026-09-01", "Manager", "Risk Management", "Ethics", 3, "Mgr", "Escalated an issue", 2, "Yes", "Manager", "Validated"]])
sheet("Projects", ["ProjectID", "ProjectTitle", "ProjectSponsor", "YLPAdviser", "Status", "CompletionPercentage", "StartDate", "TargetDate"], [["W-P1", "Workbook project", "Sponsor W", "Adviser W", "In Progress", 30, "2026-06-01", "2027-01-31"]])
sheet("ProjectParticipants", ["ProjectID", "EmployeeID", "ParticipantRole", "IndividualContribution"], [["W-P1", "W-01", "Lead", "Charter"], ["W-P1", "W-02", "Analysis", ""], ["W-P1", "W-03", "Tech", "Prototype"]])
sheet("Mentoring", ["EmployeeID", "Mentor", "MeetingDate", "DiscussionTheme", "ActionCommitted"], [["W-01", "Mentor W", "2026-08-01", "Start", "Plan"]])
sheet("Scratch notes", ["Some", "Random", "Columns"], [["a", "b", "c"], ["d", "e", "f"]])
wb.save(f"{out}/YLP_Complete_Workbook.xlsx")

print("fixtures:", len(os.listdir(out)), "files")
